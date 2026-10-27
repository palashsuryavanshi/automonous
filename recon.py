"""Excel reconciliation engine for the Automonous Automation Tools.

The module is intentionally self-contained so it can be tested without Flask.
It focuses on accounting-grade, explainable matching rather than exact equality.
"""

import datetime
import decimal
import io
import json
import os
import re
import tempfile
import unicodedata
import uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook
from rapidfuzz import fuzz


LOGICAL_FIELDS = [
    ("gstin", "GSTIN"),
    ("supplier", "Supplier name"),
    ("invoice_number", "Invoice number"),
    ("invoice_date", "Invoice date"),
    ("taxable_value", "Taxable amount"),
    ("igst", "IGST"),
    ("cgst", "CGST"),
    ("sgst", "SGST/UTGST"),
    ("cess", "Cess"),
    ("invoice_value", "Invoice value"),
    ("place_of_supply", "Place of supply"),
    ("reverse_charge", "Reverse charge"),
    ("document_type", "Document type"),
    ("financial_year", "Financial year"),
    ("reference", "Other references"),
]

HEADER_SYNONYMS = {
    "gstin": ["gstin", "gst no", "gst number", "supplier gstin", "recipient gstin", "party gstin"],
    "supplier": ["supplier", "vendor", "trade name", "party name", "supplier name", "vendor name"],
    "invoice_number": ["invoice no", "invoice number", "bill no", "bill number", "document number", "reference no", "voucher no", "doc no"],
    "invoice_date": ["invoice date", "bill date", "document date", "doc date", "date"],
    "taxable_value": ["taxable value", "taxable amount", "assessable value", "taxable"],
    "igst": ["igst", "integrated tax", "igst amount"],
    "cgst": ["cgst", "central tax", "cgst amount"],
    "sgst": ["sgst", "state tax", "utgst", "sgst amount", "utgst amount"],
    "cess": ["cess", "cess amount"],
    "invoice_value": ["invoice value", "gross amount", "total bill value", "invoice amount", "total amount", "grand total"],
    "place_of_supply": ["pos", "place of supply", "supply state", "state"],
    "reverse_charge": ["rcm", "reverse charge", "reverse charge mechanism"],
    "document_type": ["invoice type", "document type", "credit debit note", "doc type"],
    "financial_year": ["fy", "financial year"],
    "reference": ["purchase order", "po", "reference", "voucher", "ref"],
}

DEFAULT_TOLERANCE = decimal.Decimal("1.00")


def _clean_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN
        return ""
    text = str(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\xa0", " ").replace("\u200b", "")
    text = re.sub(r"[\x00-\x1f\x7f]", "", text)
    return text.strip()


def _normalize_invoice_number(value, policy="standard"):
    original = _clean_text(value)
    if not original:
        return {"original": original, "normalized": "", "tokens": [], "operations": []}

    operations = []
    text = original
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u2013", "-").replace("\u2014", "-").replace("\u2212", "-")
    text = re.sub(r"\s+", " ", text).strip()
    operations.append("trim/whitespace")

    # Standardize separators: treat / _ and spaces as dashes for tokenization.
    working = text
    working = re.sub(r"[/_\s]+", "-", working)
    operations.append("separator normalization")

    # Split into tokens on dashes.
    raw_tokens = [t for t in working.split("-") if t != ""]
    tokens = []
    for token in raw_tokens:
        # Further split camelCase/alphanumeric boundaries? Keep token as-is first.
        tokens.append(token)

    # Case normalization.
    tokens = [t.upper() for t in tokens]
    operations.append("uppercase")

    # Remove leading zeros from numeric tokens unless strict.
    if policy != "strict":
        new_tokens = []
        for token in tokens:
            if re.fullmatch(r"0+\d+", token):
                new_tokens.append(token.lstrip("0") or "0")
                operations.append("leading-zero removal")
            else:
                new_tokens.append(token)
        tokens = new_tokens

    normalized = "-".join(tokens)
    return {
        "original": original,
        "normalized": normalized,
        "tokens": tokens,
        "operations": sorted(set(operations)),
    }


def normalize_gstin(value):
    text = _clean_text(value).upper().replace(" ", "")
    return text


def gstin_valid(value):
    text = normalize_gstin(value)
    if not text:
        return False
    return bool(re.fullmatch(r"[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]", text))


def normalize_supplier(value):
    text = _clean_text(value).casefold()
    text = re.sub(r"[^\w\s]", " ", text, flags=re.UNICODE)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def parse_money(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, decimal.Decimal)):
        try:
            return decimal.Decimal(str(value))
        except Exception:
            return None
    text = _clean_text(value)
    if not text:
        return None
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative = True
        text = text[1:-1]
    text = text.replace(",", "").replace("₹", "").replace("Rs", "").replace("rs", "")
    text = text.replace("INR", "").replace("inr", "").strip()
    text = text.replace(" ", "")
    if text.startswith("-"):
        negative = True
        text = text[1:]
    if not re.fullmatch(r"\d+(\.\d+)?", text):
        return None
    try:
        result = decimal.Decimal(text)
    except Exception:
        return None
    return -result if negative else result


def parse_date(value, date_format="auto"):
    if value is None or value == "":
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial date
        try:
            base = datetime.date(1899, 12, 30)
            return base + datetime.timedelta(days=int(value))
        except Exception:
            return None
    text = _clean_text(value)
    if not text:
        return None
    # Excel serial as text
    if re.fullmatch(r"\d{4,6}", text):
        try:
            base = datetime.date(1899, 12, 30)
            return base + datetime.timedelta(days=int(text))
        except Exception:
            pass

    formats = []
    if date_format == "dmy":
        formats = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %m %Y"]
    elif date_format == "mdy":
        formats = ["%m/%d/%Y", "%m-%d-%Y", "%m.%d.%Y", "%m %d %Y"]
    elif date_format == "ymd":
        formats = ["%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d", "%Y %m %d"]
    else:
        formats = [
            "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%m-%d-%Y",
            "%d.%m.%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%b %d, %Y",
            "%B %d, %Y", "%d-%b-%Y", "%d-%B-%Y",
        ]

    for fmt in formats:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _normalize_header(text):
    text = _clean_text(text).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def detect_header_row(rows, max_scan=12):
    best_index = 0
    best_score = -1
    limit = min(len(rows), max_scan)
    for i in range(limit):
        row = rows[i] if i < len(rows) else []
        cells = [_clean_text(c) for c in row]
        non_empty = [c for c in cells if c != ""]
        if not non_empty:
            continue
        unique = len(set(non_empty))
        score = len(non_empty) + unique
        # Prefer rows with more text-like headers.
        text_like = sum(1 for c in non_empty if not re.fullmatch(r"[\d\.,\-/]+", c))
        score += text_like
        if score > best_score:
            best_score = score
            best_index = i
    return best_index


def auto_map(headers_a, headers_b):
    norm_a = {h: _normalize_header(h) for h in headers_a}
    norm_b = {h: _normalize_header(h) for h in headers_b}
    mapping = {}
    for field, label in LOGICAL_FIELDS:
        best_a = None
        best_b = None
        best_conf_a = 0.0
        best_conf_b = 0.0
        synonyms = HEADER_SYNONYMS.get(field, [])
        for h, nh in norm_a.items():
            conf = 0.0
            for syn in synonyms:
                if nh == syn:
                    conf = 1.0
                    break
                if syn in nh or nh in syn:
                    conf = max(conf, 0.75)
            # token overlap
            if conf < 0.75:
                tokens = set(nh.split())
                for syn in synonyms:
                    st = set(syn.split())
                    if tokens and st:
                        overlap = len(tokens & st) / max(len(tokens), len(st))
                        conf = max(conf, overlap * 0.6)
            if conf > best_conf_a:
                best_conf_a = conf
                best_a = h
        for h, nh in norm_b.items():
            conf = 0.0
            for syn in synonyms:
                if nh == syn:
                    conf = 1.0
                    break
                if syn in nh or nh in syn:
                    conf = max(conf, 0.75)
            if conf < 0.75:
                tokens = set(nh.split())
                for syn in synonyms:
                    st = set(syn.split())
                    if tokens and st:
                        overlap = len(tokens & st) / max(len(tokens), len(st))
                        conf = max(conf, overlap * 0.6)
            if conf > best_conf_b:
                best_conf_b = conf
                best_b = h
        # Avoid weak mappings.
        if best_conf_a < 0.5:
            best_a = None
            best_conf_a = 0.0
        if best_conf_b < 0.5:
            best_b = None
            best_conf_b = 0.0
        mapping[field] = {
            "label": label,
            "a": best_a,
            "b": best_b,
            "confidence": round((best_conf_a + best_conf_b) / 2, 2) if (best_a or best_b) else 0.0,
        }
    return mapping


def read_sheet_rows(file_path, sheet_name):
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    wb.close()
    return rows


def _file_ext(file_path):
    return Path(file_path).suffix.lower()


def get_sheet_names(file_path):
    ext = _file_ext(file_path)
    if ext == ".csv":
        return ["CSV"]
    if ext == ".xls":
        import xlrd

        wb = xlrd.open_workbook(file_path)
        return list(wb.sheet_names())
    wb = load_workbook(file_path, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def read_sheet_rows(file_path, sheet_name):
    ext = _file_ext(file_path)
    if ext == ".csv":
        import csv

        rows = []
        with open(file_path, "r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            for row in reader:
                rows.append(list(row))
        max_len = max((len(r) for r in rows), default=0)
        for r in rows:
            while len(r) < max_len:
                r.append("")
        return rows
    if ext == ".xls":
        import xlrd

        wb = xlrd.open_workbook(file_path)
        sh = wb.sheet_by_name(sheet_name)
        rows = []
        for i in range(sh.nrows):
            rows.append(list(sh.row_values(i)))
        max_len = max((len(r) for r in rows), default=0)
        for r in rows:
            while len(r) < max_len:
                r.append("")
        return rows
    wb = load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    wb.close()
    return rows


def rows_to_records(rows, header_row_index, mapping, date_format="auto", side="a"):
    """Convert raw rows into normalized records using a combined mapping."""
    if not rows:
        return [], []
    header_row = rows[header_row_index]
    headers = [_clean_text(h) for h in header_row]
    header_index = {h: i for i, h in enumerate(headers) if h != ""}

    records = []
    excluded = []
    for idx, row in enumerate(rows):
        excel_row = idx + 1
        if idx <= header_row_index:
            continue
        # Skip fully empty
        if all(_clean_text(c) == "" for c in row):
            continue
        # Detect repeated header row
        row_text = [_clean_text(c).lower() for c in row]
        if row_text == [h.lower() for h in headers]:
            excluded.append({"row": excel_row, "reason": "Repeated header row"})
            continue
        # Detect total/subtotal rows
        joined = " ".join(row_text)
        if any(k in joined for k in ["total", "grand total", "subtotal"]):
            # Only exclude if mostly empty/numeric summary
            non_empty = [c for c in row if _clean_text(c) != ""]
            if len(non_empty) <= max(3, len(headers) // 2):
                excluded.append({"row": excel_row, "reason": "Summary/total row"})
                continue

        # Build record
        def get_col(field):
            info = mapping.get(field, {})
            if isinstance(info, dict):
                return info.get(side)
            return info

        record = {"_row": excel_row, "_raw": [_clean_text(c) for c in row]}
        has_any = False
        for field, _label in LOGICAL_FIELDS:
            col = get_col(field)
            val = None
            if col and col in header_index:
                i = header_index[col]
                if i < len(row):
                    val = row[i]
            record[field] = val
            if val not in (None, ""):
                has_any = True

        # Eligibility: need invoice number or gstin or amount
        inv = _clean_text(record.get("invoice_number"))
        gstin = _clean_text(record.get("gstin"))
        amt = record.get("invoice_value") or record.get("taxable_value")
        if not (inv or gstin or amt not in (None, "")):
            excluded.append({"row": excel_row, "reason": "Incomplete/invalid row"})
            continue

        # Normalize
        inv_norm = _normalize_invoice_number(record.get("invoice_number"))
        record["_invoice_norm"] = inv_norm["normalized"]
        record["_invoice_tokens"] = inv_norm["tokens"]
        record["_gstin_norm"] = normalize_gstin(record.get("gstin"))
        record["_supplier_norm"] = normalize_supplier(record.get("supplier"))
        record["_date"] = parse_date(record.get("invoice_date"), date_format)
        record["_taxable"] = parse_money(record.get("taxable_value"))
        record["_igst"] = parse_money(record.get("igst"))
        record["_cgst"] = parse_money(record.get("cgst"))
        record["_sgst"] = parse_money(record.get("sgst"))
        record["_cess"] = parse_money(record.get("cess"))
        record["_invoice_value"] = parse_money(record.get("invoice_value"))
        records.append(record)
    return records, excluded


def _map_for_side(mapping, side):
    """Convert combined mapping {field:{a:,b:}} into {field: header} for one side."""
    result = {}
    for field, info in mapping.items():
        if isinstance(info, dict):
            result[field] = info.get(side)
        else:
            result[field] = info
    return result


def normalize_records_side(rows, header_row_index, mapping_side, date_format="auto"):
    # mapping_side: {field: header_name}
    combined = {f: {"a": h, "b": h} for f, h in mapping_side.items()}
    return rows_to_records(rows, header_row_index, combined, date_format=date_format)


def detect_duplicates(records):
    dup_groups = {}
    for r in records:
        key = (r.get("_gstin_norm", ""), r.get("_invoice_norm", ""))
        if not r.get("_invoice_norm"):
            key = ("", f"ROW-{r['_row']}")
        dup_groups.setdefault(key, []).append(r)
    duplicate_rows = set()
    for key, group in dup_groups.items():
        if len(group) > 1:
            for r in group:
                duplicate_rows.add(r["_row"])
    return duplicate_rows


def score_candidate(a, b, tolerance=DEFAULT_TOLERANCE):
    # Returns (score 0-100, details dict, hard_block bool)
    details = {}
    hard_block = False

    # Hard constraints
    ga = a.get("_gstin_norm", "")
    gb = b.get("_gstin_norm", "")
    if ga and gb and gstin_valid(ga) and gstin_valid(gb) and ga != gb:
        hard_block = True
        details["gstin"] = "Conflict"

    ia = a.get("_invoice_norm", "")
    ib = b.get("_invoice_norm", "")
    if ia and ib and ia != ib:
        hard_block = True
        details["invoice_number"] = "Mismatch"

    da = a.get("document_type")
    db = b.get("document_type")
    if da and db and _clean_text(da).lower() != _clean_text(db).lower():
        hard_block = True
        details["document_type"] = "Mismatch"

    if hard_block:
        return 0, details, True

    # Available weights
    weights = {
        "invoice_number": 40,
        "gstin": 25,
        "supplier": 10,
        "date": 10,
        "invoice_value": 10,
        "tax": 5,
    }
    available = 0
    earned = 0

    # invoice number similarity
    if ia or ib:
        available += weights["invoice_number"]
        if ia and ib:
            sim = fuzz.ratio(ia, ib) / 100.0
            earned += weights["invoice_number"] * sim
            details["invoice_number"] = f"{round(sim*100)}% similar"
        else:
            details["invoice_number"] = "Missing in one record"

    # gstin
    if ga or gb:
        available += weights["gstin"]
        if ga and gb:
            if ga == gb:
                earned += weights["gstin"]
                details["gstin"] = "Match"
            else:
                details["gstin"] = "Different"
        else:
            details["gstin"] = "Missing in one record"

    # supplier
    sa = a.get("_supplier_norm", "")
    sb = b.get("_supplier_norm", "")
    if sa or sb:
        available += weights["supplier"]
        if sa and sb:
            sim = fuzz.token_sort_ratio(sa, sb) / 100.0
            earned += weights["supplier"] * sim
            details["supplier"] = f"{round(sim*100)}% similar"
        else:
            details["supplier"] = "Missing in one record"

    # date
    da_date = a.get("_date")
    db_date = b.get("_date")
    if da_date or db_date:
        available += weights["date"]
        if da_date and db_date:
            delta = abs((da_date - db_date).days)
            if delta == 0:
                earned += weights["date"]
                details["date"] = "Same date"
            elif delta <= 7:
                earned += weights["date"] * 0.5
                details["date"] = f"{delta} days apart"
            elif delta <= 30:
                earned += weights["date"] * 0.2
                details["date"] = f"{delta} days apart"
            else:
                details["date"] = f"{delta} days apart"
        else:
            details["date"] = "Missing in one record"

    # invoice value
    va = a.get("_invoice_value")
    vb = b.get("_invoice_value")
    if va is not None or vb is not None:
        available += weights["invoice_value"]
        if va is not None and vb is not None:
            diff = abs(va - vb)
            if diff <= tolerance:
                earned += weights["invoice_value"]
                details["invoice_value"] = "Match within tolerance"
            elif va != 0 and (diff / max(abs(va), abs(vb))) <= decimal.Decimal("0.02"):
                earned += weights["invoice_value"] * 0.4
                details["invoice_value"] = f"Difference {diff}"
            else:
                details["invoice_value"] = f"Difference {diff}"
        else:
            details["invoice_value"] = "Missing in one record"

    # tax components
    tax_fields = ["_taxable", "_igst", "_cgst", "_sgst", "_cess"]
    has_tax = any(a.get(f) is not None for f in tax_fields) or any(b.get(f) is not None for f in tax_fields)
    if has_tax:
        available += weights["tax"]
        # compare total tax-ish
        def tax_total(rec):
            vals = [rec.get(f) for f in tax_fields]
            vals = [v for v in vals if v is not None]
            return sum(vals, decimal.Decimal("0")) if vals else None
        ta = tax_total(a)
        tb = tax_total(b)
        if ta is not None and tb is not None:
            diff = abs(ta - tb)
            if diff <= tolerance:
                earned += weights["tax"]
                details["tax"] = "Match within tolerance"
            else:
                details["tax"] = f"Difference {diff}"
        else:
            details["tax"] = "Partial tax data"

    if available == 0:
        return 0, details, False
    score = int(round((earned / available) * 100))
    return score, details, False


def discrepancy_list(a, b, tolerance=DEFAULT_TOLERANCE):
    diffs = []
    field_map = {
        "invoice_number": ("Invoice number", "_invoice_norm", None),
        "gstin": ("GSTIN", "_gstin_norm", None),
        "supplier": ("Supplier", "_supplier_norm", None),
        "invoice_date": ("Invoice date", "_date", None),
        "taxable_value": ("Taxable value", "_taxable", "money"),
        "igst": ("IGST", "_igst", "money"),
        "cgst": ("CGST", "_cgst", "money"),
        "sgst": ("SGST/UTGST", "_sgst", "money"),
        "cess": ("Cess", "_cess", "money"),
        "invoice_value": ("Invoice value", "_invoice_value", "money"),
    }
    for field, (label, key, kind) in field_map.items():
        va = a.get(key)
        vb = b.get(key)
        if va is None and vb is None:
            continue
        if kind == "money":
            if va is not None and vb is not None:
                if abs(va - vb) > tolerance:
                    diffs.append(f"{label}: {va} vs {vb} (diff {abs(va-vb)})")
            elif va != vb:
                diffs.append(f"{label}: {va} vs {vb}")
        else:
            if va != vb:
                diffs.append(f"{label}: {va} vs {vb}")
    return diffs


def reconcile(file_a_path, sheet_a, header_a, file_b_path, sheet_b, header_b,
             mapping, date_format="auto", tolerance=DEFAULT_TOLERANCE,
             review_threshold=70, auto_threshold=85):
    rows_a = read_sheet_rows(file_a_path, sheet_a)
    rows_b = read_sheet_rows(file_b_path, sheet_b)

    records_a, excluded_a = rows_to_records(rows_a, header_a, mapping, date_format=date_format, side="a")
    records_b, excluded_b = rows_to_records(rows_b, header_b, mapping, date_format=date_format, side="b")

    # Mark duplicates
    dup_a = detect_duplicates(records_a)
    dup_b = detect_duplicates(records_b)
    for r in records_a:
        r["_duplicate"] = r["_row"] in dup_a
    for r in records_b:
        r["_duplicate"] = r["_row"] in dup_b

    # Candidate generation
    candidates_by_a = {}
    for a in records_a:
        cands = []
        for b in records_b:
            score, details, blocked = score_candidate(a, b, tolerance=tolerance)
            if blocked:
                continue
            if score >= 40:
                cands.append({"b": b, "score": score, "details": details})
        cands.sort(key=lambda x: x["score"], reverse=True)
        candidates_by_a[id(a)] = cands

    # One-to-one assignment
    assigned_b = set()
    results = []
    for a in records_a:
        cands = candidates_by_a[id(a)]
        if not cands:
            status = "Missing in File B"
            best = None
        else:
            best = cands[0]
            # ambiguity: second candidate close
            ambiguous = len(cands) > 1 and (best["score"] - cands[1]["score"] < 8)
            if best["score"] >= auto_threshold and not ambiguous and id(best["b"]) not in assigned_b:
                assigned_b.add(id(best["b"]))
                diffs = discrepancy_list(a, best["b"], tolerance=tolerance)
                if diffs:
                    status = "Matched with discrepancies"
                else:
                    status = "Matched"
            elif best["score"] >= review_threshold:
                status = "Needs Review"
            else:
                status = "Missing in File B"
        results.append({
            "a": a,
            "b": best["b"] if best and status not in ("Missing in File B",) else None,
            "score": best["score"] if best else 0,
            "details": best["details"] if best else {},
            "status": status,
            "discrepancies": discrepancy_list(a, best["b"], tolerance=tolerance) if (best and best["b"] is not None) else [],
            "candidate_count": len(cands),
        })

    # Missing in File A: B records not assigned
    assigned_b_ids = {id(r["b"]) for r in results if r["b"] is not None}
    for b in records_b:
        if id(b) not in assigned_b_ids:
            results.append({
                "a": None,
                "b": b,
                "score": 0,
                "details": {},
                "status": "Missing in File A",
                "discrepancies": [],
                "candidate_count": 0,
            })

    # Duplicate status override
    for r in results:
        a = r["a"]
        b = r["b"]
        if a is not None and a.get("_duplicate"):
            r["status"] = "Duplicates" if r["status"] not in ("Missing in File B",) else r["status"]
        if b is not None and b.get("_duplicate"):
            if r["status"] not in ("Missing in File A",):
                r["status"] = "Duplicates"

    # Summary
    total_a = len(records_a)
    total_b = len(records_b)
    matched = sum(1 for r in results if r["status"] == "Matched")
    matched_disc = sum(1 for r in results if r["status"] == "Matched with discrepancies")
    missing_a = sum(1 for r in results if r["status"] == "Missing in File A")
    missing_b = sum(1 for r in results if r["status"] == "Missing in File B")
    review = sum(1 for r in results if r["status"] == "Needs Review")
    duplicates = sum(1 for r in results if r["status"] == "Duplicates")
    unique_pairs = matched + matched_disc
    denom = max(total_a, total_b) if max(total_a, total_b) else 1
    recon_pct = round(unique_pairs / denom * 100, 2)

    summary = {
        "total_a": total_a,
        "total_b": total_b,
        "matched": matched,
        "matched_with_discrepancies": matched_disc,
        "missing_in_a": missing_a,
        "missing_in_b": missing_b,
        "needs_review": review,
        "duplicates": duplicates,
        "reconciliation_percentage": recon_pct,
    }

    return {
        "summary": summary,
        "results": results,
        "excluded_a": excluded_a,
        "excluded_b": excluded_b,
        "records_a": records_a,
        "records_b": records_b,
    }


def create_report(recon_result, output_path):
    wb = Workbook()
    # Remove default
    wb.remove(wb.active)

    summary = recon_result["summary"]
    results = recon_result["results"]

    def style_header(ws, headers):
        ws.append(headers)
        for c in ws[1]:
            c.font = c.font.copy(bold=True)
        ws.freeze_panes = "A2"

    # Summary
    ws = wb.create_sheet("Summary")
    ws.append(["Metric", "Value"])
    for k, v in summary.items():
        ws.append([k, v])

    # Helper to write result rows
    def result_rows(status_filter=None):
        out = []
        for r in results:
            if status_filter and r["status"] != status_filter:
                continue
            a = r["a"] or {}
            b = r["b"] or {}
            out.append([
                a.get("invoice_number", ""),
                b.get("invoice_number", ""),
                a.get("_invoice_norm", ""),
                b.get("_invoice_norm", ""),
                a.get("gstin", ""),
                b.get("gstin", ""),
                a.get("supplier", ""),
                b.get("supplier", ""),
                a.get("invoice_date", ""),
                b.get("invoice_date", ""),
                a.get("taxable_value", ""),
                b.get("taxable_value", ""),
                a.get("igst", ""),
                b.get("igst", ""),
                a.get("cgst", ""),
                b.get("cgst", ""),
                a.get("sgst", ""),
                b.get("sgst", ""),
                a.get("cess", ""),
                b.get("cess", ""),
                a.get("invoice_value", ""),
                b.get("invoice_value", ""),
                a.get("_row", ""),
                b.get("_row", ""),
                r["score"],
                r["status"],
                "; ".join(r["discrepancies"]),
            ])
        return out

    headers = [
        "File A Invoice", "File B Invoice", "Norm A", "Norm B",
        "GSTIN A", "GSTIN B", "Supplier A", "Supplier B",
        "Date A", "Date B", "Taxable A", "Taxable B",
        "IGST A", "IGST B", "CGST A", "CGST B", "SGST A", "SGST B",
        "Cess A", "Cess B", "Value A", "Value B",
        "Row A", "Row B", "Score", "Status", "Discrepancies",
    ]

    def write_sheet(name, status_filter=None):
        ws = wb.create_sheet(name)
        style_header(ws, headers)
        for row in result_rows(status_filter):
            ws.append(row)

    write_sheet("Matched", "Matched")
    write_sheet("Discrepancies", "Matched with discrepancies")
    # Missing in File A
    ws = wb.create_sheet("Missing in File A")
    style_header(ws, headers)
    for row in result_rows("Missing in File A"):
        ws.append(row)
    ws = wb.create_sheet("Missing in File B")
    style_header(ws, headers)
    for row in result_rows("Missing in File B"):
        ws.append(row)
    ws = wb.create_sheet("Needs Review")
    style_header(ws, headers)
    for row in result_rows("Needs Review"):
        ws.append(row)
    ws = wb.create_sheet("Duplicates")
    style_header(ws, headers)
    for row in result_rows("Duplicates"):
        ws.append(row)
    ws = wb.create_sheet("All Records")
    style_header(ws, headers)
    for row in result_rows(None):
        ws.append(row)

    # Field Mapping
    ws = wb.create_sheet("Field Mapping")
    ws.append(["Logical field", "File A column", "File B column"])
    for field, label in LOGICAL_FIELDS:
        pass
    # mapping not passed here; skip detailed

    # Excluded rows
    ws = wb.create_sheet("Excluded Rows")
    ws.append(["File", "Row", "Reason"])
    for e in recon_result["excluded_a"]:
        ws.append(["File A", e["row"], e["reason"]])
    for e in recon_result["excluded_b"]:
        ws.append(["File B", e["row"], e["reason"]])

    # Audit log
    ws = wb.create_sheet("Audit Log")
    ws.append(["Timestamp", "Action"])
    ws.append([datetime.datetime.now().isoformat(), "Reconciliation completed"])

    wb.save(output_path)
    return output_path
