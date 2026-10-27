"""PDF helpers for Automonous utilities.

Currently provides OCR-based PDF -> DOCX conversion. The flow is:

1. Render every PDF page to a high-resolution image.
2. Run Tesseract OCR on each page and rebuild a searchable ("OCR'd") PDF.
3. Extract the OCR text from that PDF and write it into a .docx file.

Everything runs locally; no uploads leave the machine.
"""

import io
import os
import shutil
from pathlib import Path

import pymupdf
import pytesseract
from docx import Document
from docx.shared import Pt
from PIL import Image
from pypdf import PdfReader, PdfWriter


class OcrError(RuntimeError):
    """Raised when OCR or conversion cannot be completed."""


def _find_tesseract():
    """Return the Tesseract executable path or raise a helpful error."""
    env_path = os.environ.get("TESSERACT_CMD")
    candidates = []

    if env_path:
        candidates.append(Path(env_path))

    candidates.extend(
        [
            Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
            Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
        ]
    )

    local_app = os.environ.get("LOCALAPPDATA")
    if local_app:
        candidates.append(Path(local_app) / "Programs" / "Tesseract-OCR" / "tesseract.exe")

    found = shutil.which("tesseract")
    if found:
        candidates.append(Path(found))

    for candidate in candidates:
        if candidate and candidate.is_file():
            return candidate

    raise OcrError(
        "Tesseract OCR was not found. Install it from "
        "https://github.com/UB-Mannheim/tesseract/wiki and restart Automonous."
    )


def _page_images(pdf_bytes, dpi=300):
    """Yield rendered page images from a PDF."""
    document = pymupdf.open(stream=pdf_bytes, filetype="pdf")

    try:
        for page in document:
            pixmap = page.get_pixmap(dpi=dpi)
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            yield image
    finally:
        document.close()


def _build_ocr_pdf(pdf_bytes, lang="eng"):
    """OCR each page and return a searchable PDF as bytes."""
    pytesseract.pytesseract.tesseract_cmd = str(_find_tesseract())

    writer = PdfWriter()
    pages = 0

    for image in _page_images(pdf_bytes):
        page_pdf = pytesseract.image_to_pdf_or_hocr(
            image,
            extension="pdf",
            lang=lang,
        )
        reader = PdfReader(io.BytesIO(page_pdf))
        for page in reader.pages:
            writer.add_page(page)
        pages += 1

    if pages == 0:
        raise OcrError("The PDF has no pages to convert.")

    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _ocr_pdf_to_docx(ocr_pdf_bytes):
    """Extract text from an OCR'd PDF and return a .docx as bytes."""
    document = pymupdf.open(stream=ocr_pdf_bytes, filetype="pdf")
    docx = Document()

    style = docx.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    try:
        for index, page in enumerate(document):
            text = page.get_text("text").strip()

            if index > 0:
                docx.add_page_break()

            if not text:
                docx.add_paragraph("")
                continue

            for block in text.split("\n\n"):
                cleaned = " ".join(line.strip() for line in block.splitlines() if line.strip())
                if cleaned:
                    docx.add_paragraph(cleaned)
    finally:
        document.close()

    output = io.BytesIO()
    docx.save(output)
    return output.getvalue()


def pdf_to_docx(pdf_bytes, lang="eng"):
    """OCR a PDF and return (docx_bytes, ocr_pdf_bytes)."""
    if not pdf_bytes:
        raise OcrError("The uploaded PDF is empty.")

    try:
        ocr_pdf = _build_ocr_pdf(pdf_bytes, lang=lang)
    except OcrError:
        raise
    except Exception as exc:
        raise OcrError(f"OCR failed: {exc}") from exc

    try:
        docx_bytes = _ocr_pdf_to_docx(ocr_pdf)
    except Exception as exc:
        raise OcrError(f"Could not create the Word document: {exc}") from exc

    return docx_bytes, ocr_pdf